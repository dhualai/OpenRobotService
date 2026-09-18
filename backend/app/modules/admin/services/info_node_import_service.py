"""项目信息树「文件导入（AI 识别）」服务 —— 上传文档，大模型抽取信息并与现有节点匹配预览。

用途：项目信息管理 → 编辑项目信息 → 文件导入。用户上传 Word / Markdown / Excel / 文本
需求文档，由大模型（与「我要摇人」同一客户端：backend/.env 的 LLM_API_KEY，默认
DeepSeek flash，即 settings.LLM_MODEL_NAME）抽取「信息条目」，再与项目现有信息节点匹配，
按三类返回预览供前端勾选确认（本服务只读不写，落库由前端确认后逐节点走既有 CRUD 接口）：

  1. fill      —— 识别到且节点当前为空 → 勾选后直接填入；
  2. overwrite —— 识别到且节点已有不同内容 → 勾选后覆盖（前端显示 原内容 → 新内容）；
  3. unmatched —— 与文件有关但系统没有对应节点 → 勾选后作为新节点创建（附建议归属）。

匹配规则（与需求一致）：条目与节点标题「精确一致，或相似度 ≥ 0.9（满分 1）」（difflib
序列相似度，规范化空白/标点后比较）；下拉节点还要求识别值命中其可选项，否则按未匹配处理。
大模型返回的 nodeTitle 只是提示，最终匹配以后端确定性算法为准。
"""
from __future__ import annotations

import asyncio
import difflib
import io
import json
import re
import zipfile
from typing import Any, Dict, List, Optional, Tuple

from app.core.config import settings
from app.modules.admin.services.info_node_service import info_node_service

# ── 常量 ──────────────────────────────────────────────────────────

# 支持的扩展名（.doc/.xls 旧二进制格式不支持，提示另存为新格式）
TEXT_EXTENSIONS = {".md", ".markdown", ".txt", ".csv"}
ALLOWED_EXTENSIONS = TEXT_EXTENSIONS | {".docx", ".xlsx"}

MAX_FILE_BYTES = 10 * 1024 * 1024  # 上传文件上限 10MB
MAX_TEXT_CHARS = 100_000           # 送大模型的正文上限（超出截断）
MAX_SHEET_ROWS = 2000              # Excel 每个工作表最多读取行数
SIMILARITY_THRESHOLD = 0.9         # 节点名相似度阈值（满分 1，与需求一致）
NAME_MISMATCH_THRESHOLD = 0.6      # 项目名一致性阈值：低于该相似度且互不包含 → 判定不一致（提醒可能导错文件）
MAX_NODE_DEPTH = 4                 # 信息树最大层级（与前端 PROJECT_INFO_MAX_DEPTH 一致）
LLM_TEMPERATURE = 0.2              # 抽取任务用低温度，减少发散
LLM_TIMEOUT_SECONDS = 120.0

SYSTEM_PROMPT = "你是项目信息整理助手，只输出 JSON，不输出任何解释文字或 Markdown 代码块。"

# AGV 车型目录（66 款）——信息树里存在车辆/车型节点时注入提示词，规范车型写法。
# 型号清单与前端 frontend/src/shared/utils/vehicleModels.ts 的 VEHICLE_MODEL_CODES
# **必须一字不差**，改动时两边同步；库里车型1/车型2 的 config.options 由
# 迁移 4a7c2e9d1b53 同步，别只改代码。
VEHICLE_MODEL_CODES: List[str] = [
    'UHX-01', 'EXP15', 'RPG201', 'XCART', 'XC1031', 'XC1051', 'XC1061', 'XCS101U',
    'XCU0051', 'XCL0051', 'XCO0051', 'XCB031', 'XCF101', 'XFC001', 'XFC002', 'XCD0051',
    'XCD031', 'XCD061', 'XCD062', 'XCD101', 'XCD151', 'XCD202Y', 'XCD301', 'XCD501',
    'XCT201', 'XTD401', 'XTD601', 'XPA152', 'XPC151', 'XPG151', 'XP1151', 'XP1152',
    'XP1153', 'XP1201', 'XP3201', 'XPL201', 'XPL201P', 'XPL201T', 'XPL301', 'XPL501',
    'XQE122', 'XQE151', 'XQC161', 'XQC163', 'XQC201', 'XQS151', 'XQS181', 'XS1151',
    'XS1152', 'XS1201', 'XS2201', 'XSC081', 'XSC121', 'XSC151', 'XSC201', 'XSF101',
    'XSG121', 'XNA101', 'XNA121', 'XNA151', 'XFL151E', 'XFL201', 'XFL301', 'XFL351',
    'XORD1', 'XORD3',
]

# 旧型号名 → 现型号（文件里可能还写着老名字，别让它因为改名而匹配不上下拉）。
# 只有改过号的才登记；比对时两边都过 _norm，所以大小写/连字符差异不用写在这里。
VEHICLE_MODEL_ALIASES: Dict[str, str] = {
    'XS1161': 'XS1201',
}

# 型号 → 中文全称，只给提示词用（文件里写「潜伏顶升搬运机器人 1500kg」时靠它对上 XCD151）。
# 没登记名称的型号在提示词里只出现型号本身——宁可不写，也别把猜测的载重/结构写进去。
VEHICLE_MODEL_NAMES: Dict[str, str] = {
    'XC1051': '背负式搬运机器人 500 kg',
    'XC1061': '跟随潜伏式机器人 600 kg',
    'XCD031': '潜伏顶升搬运机器人 300 kg',
    'XCD061': '潜伏顶升搬运机器人 600 kg',
    'XCD101': '潜伏顶升搬运机器人 1000 kg',
    'XCD151': '潜伏顶升搬运机器人 1500 kg',
    'XCD301': '全向潜伏顶升式机器人 3000 kg',
    'XCD501': '重载潜伏顶升式机器人 5000 kg',
    'EXP15': '极简自动搬运车 1500 kg',
    'RPG201': '踏板式自动搬运车 2000 kg',
    'XPC151': '极简智能搬运车 1500 kg',
    'XPG151': '步行式自动搬运车 1500 kg',
    'XSG121': '堆高式自动搬运车 1200 kg',
    'XCF101': '潜伏式叉车机器人 1000 kg',
    'XP1151': '点对点智能搬运机器人 1500 kg',
    'XP1152': '点对点智能搬运机器人 1500 kg',
    'XP1201': '薄背搬运式机器人 2000 kg',
    'XP3201': '室内外多场景智能搬运机器人 2000 kg',
    'XPL201': '高速重载智能搬运机器人 2000 kg',
    'XPL201P': '物流专用高速搬运机器人 2000 kg',
    'XPL201T': '薄背物流专用搬运机器人 2000 kg',
    'XPL301': '高速重载智能搬运机器人 3000 kg',
    'XPL501': '高速重载智能搬运机器人 5000 kg',
    'XFL201': '平衡重式机器人 2000 kg',
    'XNA101': '双侧叉平衡重式机器人 1000 kg',
    'XNA121': '双侧叉平衡重式机器人 1200 kg',
    'XNA151': '单侧叉平衡重式机器人 1500 kg',
    'XQE151': '平衡重式机器人 1500 kg',
    'XS1151': '薄背堆高机器人 1500 kg',
    'XS1152': '薄背堆高机器人 1500 kg',
    'XS1201': '超薄托盘堆垛机器人 2000 kg',   # 原 XS1161 改号（产品资料：2.0 吨超薄托盘堆垛）
    'XS2201': '重载堆高机器人 2000 kg',
    'XSC081': '平衡重式堆高机器人 800 kg',
    'XSC121': '平衡重式堆高机器人 1200 kg',
    'XSC151': '平衡重式堆高机器人 1500 kg',
    'XSC201': '平衡重式堆高机器人 2000 kg',
    'XSF101': '单侧叉堆高式机器人 1000 kg',
    'XQC161': '室内前移式机器人 1600 kg',
    'XQC201': '室内前移式机器人 2000 kg',
    'XQE122': '室内前移式机器人 1200 kg',
    'XQS151': '室外前移式机器人 1500 kg',
    'XQS181': '室外前移式机器人 1800 kg',
    'XCART': '智能观光车 500 kg',
    'XCT201': '室内牵引式机器人 2000 kg',
    'XTD401': '室外牵引式机器人 4000 kg',
    'XTD601': '室外牵引式机器人 6000 kg',
    'XCU0051': '料箱存取机器人 50/50+50×4 kg',
    'XCB031': '单臂具身机器人 背负 300 kg / 抓取 2-5 kg',
    'XCL0051': '料箱转运具身机器人 50×4 kg',
    'XCO0051': '料箱拣选具身机器人 5/50+50 kg',
    'XFC001': '数智飞仓',
    'XFC002': '数智飞仓',
}

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# 相似度比较前抹掉的空白与标点（中英文标点、括号、分隔符、连接符等）
_NORM_STRIP_RE = re.compile(r"[\s　()（）\[\]【】<>《》\"'“”‘’,，。.：:；;、·|/\\\-—_~]+")
# 从识别值里抠型号用：字母/数字/连字符组成的片段（「XCD101 潜伏顶升搬运机器人」→ XCD101）
_MODEL_TOKEN_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9\-]*")


# ── 小工具 ────────────────────────────────────────────────────────

def _norm(text: Any) -> str:
    """规范化标题/内容：去空白与常见标点、统一小写，用于精确比较。"""
    return _NORM_STRIP_RE.sub("", str(text or "")).lower()


def _ratio(a: str, b: str) -> float:
    """序列相似度（满分 1）。两边都为空串时视为 1。"""
    if not a and not b:
        return 1.0
    return difflib.SequenceMatcher(None, a, b).ratio()


def _get_project_name(project_id: str) -> str:
    """项目 ID（= project.id/code）→ 项目名称；查不到返回空串。"""
    from app.modules.admin.models_das.models import Project
    from app.modules.admin.services.info_node_service import SessionLocal

    db = SessionLocal()
    try:
        row = db.query(Project).filter(Project.id == project_id).first()
        return (row.name or "") if row else ""
    finally:
        db.close()


def project_name_mismatch(system_name: str, file_name: Optional[str]) -> bool:
    """文件中的项目名与系统内项目名是否「确实不一致」（用于提醒可能导错文件）。

    - 文件里没识别到项目名（None/空）→ False（无从比较，不打扰用户）；
    - 规范化后相等、或一方包含另一方（如「中力越南项目」vs「中力越南项目（二期）」）→ False；
    - 其余情况按序列相似度 < 0.6 判定为不一致。
    """
    a, b = _norm(system_name), _norm(file_name or "")
    if not a or not b:
        return False
    if a == b or a in b or b in a:
        return False
    return _ratio(a, b) < NAME_MISMATCH_THRESHOLD


def _select_state(node: Dict) -> Tuple[str, List[str]]:
    """下拉节点的 (selected, options)；坏数据按空处理。

    value 有两种形态：接口树里已是 {selected, options} 字典（_encode_value 现拼的），
    旧调用方（如导入源）给的是同结构 JSON 字符串——两种都认。只认字符串会让所有下拉
    在这里退化成空：选项空 → 值匹配不上、当前值读成空串。
    """
    raw = node.get("value")
    parsed: Any = None
    if isinstance(raw, dict):
        parsed = raw
    elif isinstance(raw, str):
        try:
            parsed = json.loads(raw)
        except (json.JSONDecodeError, TypeError):
            parsed = None
    if not isinstance(parsed, dict):
        return "", []
    selected = parsed.get("selected") if isinstance(parsed.get("selected"), str) else ""
    options = [o for o in parsed.get("options", []) if isinstance(o, str)] if isinstance(parsed.get("options"), list) else []
    return selected, options


def _current_text(node: Dict) -> str:
    """节点当前内容（text 原样字符串；select 取 selected）。"""
    if node.get("content_type") == "select":
        return _select_state(node)[0]
    raw = node.get("value")
    return raw if isinstance(raw, str) else ""


# ── 第一步：按扩展名抽取纯文本 ─────────────────────────────────────

def _decode_plain_text(data: bytes) -> str:
    """文本类文件解码：UTF-8（含 BOM）优先，其次 GBK（中文文档常见），最后容错替换。"""
    for encoding in ("utf-8-sig", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _paragraph_text(paragraph) -> str:
    """取 Word 段落文字（制表符/换行按空格处理，保留单元格内可读性）。"""
    parts: List[str] = []
    for node in paragraph.iter():
        if node.tag == f"{_W_NS}t":
            parts.append(node.text or "")
        elif node.tag in (f"{_W_NS}tab", f"{_W_NS}br", f"{_W_NS}cr"):
            parts.append(" ")
    return "".join(parts).strip()


def _extract_docx_text(data: bytes) -> str:
    """抽取 .docx 正文（标准库实现：docx 本质是 zip + word/document.xml）。

    段落逐行输出；表格逐行输出（单元格用 " | " 分隔），保证「标签 | 内容」类
    两列表格仍能读出对应关系。不解析页眉页脚与图片。
    """
    import xml.etree.ElementTree as ET

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml_bytes = archive.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError("不是有效的 .docx 文件（旧版 .doc 请先另存为 .docx 再导入）") from exc

    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError as exc:
        raise ValueError("Word 文档内容解析失败，请检查文件是否损坏") from exc

    body = root.find(f"{_W_NS}body")
    if body is None:
        return ""

    lines: List[str] = []
    for child in body:
        if child.tag == f"{_W_NS}p":
            text = _paragraph_text(child)
            if text:
                lines.append(text)
        elif child.tag == f"{_W_NS}tbl":
            for row in child.findall(f"{_W_NS}tr"):
                cells = [" ".join(_paragraph_text(p) for p in cell.findall(f"{_W_NS}p")).strip()
                         for cell in row.findall(f"{_W_NS}tc")]
                if any(cells):
                    lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_xlsx_text(data: bytes) -> str:
    """抽取 .xlsx 全部工作表内容（每行制表符分隔，工作表名作小标题）。"""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 对坏文件抛出的异常类型不稳定，统一转 ValueError
        raise ValueError("不是有效的 .xlsx 文件（旧版 .xls 请先另存为 .xlsx 再导入）") from exc

    chunks: List[str] = []
    try:
        for sheet in workbook.worksheets:
            lines: List[str] = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_SHEET_ROWS:
                    lines.append("…（该工作表内容过长，已截断）")
                    break
                cells = ["" if cell is None else str(cell).strip() for cell in row]
                if any(cells):
                    lines.append("\t".join(cells).rstrip())
            if lines:
                chunks.append(f"# 工作表：{sheet.title}\n" + "\n".join(lines))
    finally:
        workbook.close()
    return "\n\n".join(chunks)


def extract_text(filename: str, data: bytes) -> str:
    """按扩展名抽取纯文本。不支持的类型 / 解析失败抛 ValueError（接口转 400）。"""
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext in (".doc", ".xls"):
        raise ValueError(f"暂不支持旧版 {ext} 格式，请先在 Office 中另存为 {ext}x 再导入")
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("仅支持 Word（.docx）、Markdown（.md）、文本（.txt/.csv）、Excel（.xlsx）文件")

    if ext == ".docx":
        text = _extract_docx_text(data)
    elif ext == ".xlsx":
        text = _extract_xlsx_text(data)
    else:
        text = _decode_plain_text(data)

    text = text.strip()
    if not text:
        raise ValueError("没有从文件中读到文字内容")
    return text


# ── 第二步：拼节点清单 + 提示词 ────────────────────────────────────

def flatten_tree(roots: List[Dict]) -> List[Dict]:
    """递归树 → 带层级路径的平铺列表（保序）。depth 从 1 起。"""
    flat: List[Dict] = []

    def walk(items: List[Dict], parent_id: Optional[str], parent_path: List[str], depth: int) -> None:
        for item in items:
            path_titles = parent_path + [item.get("title", "")]
            flat.append({
                "id": item["id"],
                "parent_id": parent_id,
                "title": item.get("title", ""),
                "content_type": item.get("content_type", "text"),
                "value": item.get("value"),
                "options": _select_state(item)[1] if item.get("content_type") == "select" else [],
                "depth": depth,
                "path": " / ".join(path_titles),
                "path_titles": path_titles,
                "has_children": bool(item.get("children")),
            })
            walk(item.get("children") or [], item["id"], path_titles, depth + 1)

    walk(roots, None, [], 1)
    return flat


def is_fillable_node(node: Dict) -> bool:
    """该节点自己能不能填值：末级字段，或本身带值类型的分组。

    与前端渲染规则、模板校验同一口径——纯 text 的非末级节点只是分组，自己没有值；
    而「车型1」（下拉 + 「数量」子节点）两者兼具，既在下拉里选型号、又要能匹配到。
    """
    return (not node["has_children"]) or (node.get("content_type") or "text") != "text"


def build_node_catalog(flat: List[Dict]) -> str:
    """给大模型看的节点清单：每行「节点路径<TAB>内容类型[<TAB>可选项]」，可填值的另标注。"""
    lines: List[str] = []
    for node in flat:
        fill_mark = "\t(可填)" if is_fillable_node(node) else ""
        options = f"\t可选项：{'|'.join(node['options'])}" if node["options"] else ""
        lines.append(f"{node['path']}\t{node['content_type']}{fill_mark}{options}")
    return "\n".join(lines)


def find_vehicle_parent_path(flat: List[Dict]) -> Optional[str]:
    """车型信息的建议归属路径：优先取「车型N」节点的父级（车辆），否则取「车辆」节点本身。

    返回 None 表示信息树里没有车辆/车型节点（提示词不注入车型清单）。
    """
    by_id = {n["id"]: n for n in flat}
    fallback: Optional[str] = None
    for node in flat:
        title = (node.get("title") or "").strip()
        if title.startswith("车型"):
            parent = by_id.get(node.get("parent_id"))
            if parent is not None:
                return parent["path"]
        elif "车辆" in title and fallback is None:
            fallback = node["path"]
    return fallback


def build_vehicle_model_catalog() -> str:
    """AGV 车型清单（「型号（中文全称）」，没登记名称的只写型号）——仅信息树有车辆/车型节点时进提示词。"""
    return "、".join(
        f"{code}（{VEHICLE_MODEL_NAMES[code]}）" if code in VEHICLE_MODEL_NAMES else code
        for code in VEHICLE_MODEL_CODES
    )


def build_import_prompt(
    catalog: str,
    file_text: str,
    project_name: str = "",
    vehicle_parent_path: Optional[str] = None,
) -> str:
    """文件识别提示词：约束大模型只抽真实信息、按 0.9 把握匹配节点、固定 JSON 输出。

    同时把「本次导入的目标项目名称」告知大模型，并要求它回传「文件里自己写的项目名称」
    （projectName），供后端比对、提醒用户可能导错了文件。
    信息树里有车辆/车型节点时（vehicle_parent_path 非空）额外注入 AGV 车型清单，
    让车型落到「车型N」下拉框的值上（与模板页「填入车型目录」同一份目录），
    数量随同一条的 quantity 字段给出，由 match_items 填进该车型下的「数量」子节点。
    """
    target = project_name or "（未提供）"
    rules = """1. 只抽取文件里明确写到的信息，禁止编造、外推或用常识补全；文件里没写的节点不要出现在结果里。
2. 每条信息给出简洁准确的内容 value（保留型号、IP、端口、日期、数量等原文细节），不超过 200 字。
3. 节点匹配从严：只有当该信息与某个「(可填)」节点语义一致、把握 ≥ 0.9（满分 1）时，nodeTitle 才填该节点的标题（逐字复制清单里的文字，不含路径）；把握不足时 nodeTitle 填 null，并在 suggestedParentPath 里给出建议归属的节点路径（从清单里选最贴切的一级/二级路径）。
4. 若匹配到内容类型为 select 的节点，value 必须是该节点可选项中的某一项（逐字一致）；没有合适选项就按未匹配处理（nodeTitle 填 null）。
5. 同一个节点最多匹配一条信息；一条信息最多匹配一个节点。
6. 与项目无关或零散无法归类的信息不要输出。
7. 另外核对：文件里如果明确写了它自己所属的项目名称（如标题、封面、表头、「项目名称」栏），把它逐字填到 projectName；文件里没写就填 null。不要照抄上面给出的目标项目名称。"""
    vehicle_block = ""
    if vehicle_parent_path:
        rules += f"""
8. 文件中提到的 AGV 车型（设备型号）属于车辆信息，一个车型输出**一条**信息：
   - 车型型号写进**车型节点的值**，不是写进标题：该车型若对应清单里「{vehicle_parent_path}」下某个「车型N」节点（内容类型 select），把 nodeTitle 填成该节点标题、value 填**清单里的车型型号**（文件写法不同时用清单写法，如「XC1051」，旧型号 XS1161 一律写 XS1201）；
   - 该车型的数量放进同一条的 quantity 字段（如「6 台」），后端会填到该车型节点下的「数量」子节点；文件没写数量就填 null。不要再单独输出「数量」条目；
   - 清单里的「车型N」节点数不够（车型比节点多）时，多出来的车型 nodeTitle 填 null、suggestedParentPath 填「{vehicle_parent_path}」，title 填清单里的车型型号，quantity 照填。"""
        vehicle_block = f"""
下面是 AGV 车型清单（「型号（中文全称）」，没带名称的型号只有型号本身），用于规范车型信息的写法：
<<<车型清单
{build_vehicle_model_catalog()}
>>>
"""
    return f"""本次导入的目标项目名称是：「{target}」（仅供你参考上下文，不要假设文件一定属于该项目）。

下面是一个项目的「信息节点清单」（每行：节点路径<TAB>内容类型[<TAB>可选项]，“/”表示层级，标了「(可填)」的节点才能填内容）：
<<<节点清单
{catalog}
>>>
{vehicle_block}
下面是待导入文件的内容：
<<<文件内容
{file_text}
>>>

请从文件内容中抽取与该项目有关的信息，并尽量对应到上面的节点清单。要求：
{rules}

只输出如下 JSON（不要输出 Markdown 代码块或任何解释；没有识别到信息时 items 为空数组）：
{{"projectName": "文件中出现的项目名称，没有则 null", "items": [{{"title": "信息名称（尽量用清单中的节点标题）", "value": "信息内容", "nodeTitle": "匹配到的节点标题，或 null", "quantity": "车型数量（只有车型条目填，其它一律 null）", "suggestedParentPath": "建议归属的节点路径，或 null"}}]}}"""


# ── 第三步：解析大模型输出 ─────────────────────────────────────────

def _load_llm_json(content: str) -> Any:
    """大模型输出 → JSON 对象（容错剥离 ```json 围栏与前后杂文字）。"""
    raw = (content or "").strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```[a-zA-Z]*\s*", "", raw)
        raw = re.sub(r"```\s*$", "", raw).strip()

    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start < 0 or end <= start:
            raise ValueError("大模型返回内容无法解析为 JSON")
        try:
            return json.loads(raw[start:end + 1])
        except json.JSONDecodeError as exc:
            raise ValueError("大模型返回内容无法解析为 JSON") from exc


def _normalize_items(items: Any) -> List[Dict[str, Optional[str]]]:
    """items 原始数组 → 归一化条目列表（字段缺失/空值过滤）。"""
    if not isinstance(items, list):
        return []
    normalized: List[Dict[str, Optional[str]]] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title") or "").strip()
        value = item.get("value")
        value = str(value).strip() if value is not None else ""
        node_title = item.get("nodeTitle")
        node_title = str(node_title).strip() if isinstance(node_title, str) and node_title.strip() else None
        parent_path = item.get("suggestedParentPath")
        parent_path = str(parent_path).strip() if isinstance(parent_path, str) and parent_path.strip() else None
        # 车型条目自带的「数量」（如「6 台」）：匹配后落进该车型节点的「数量」子节点
        quantity = item.get("quantity")
        quantity = str(quantity).strip() if quantity is not None else ""
        if not value or not (title or node_title):
            continue
        normalized.append({
            "title": title,
            "value": value,
            "node_title": node_title,
            "quantity": quantity or None,
            "suggested_parent_path": parent_path,
        })
    return normalized


def parse_llm_payload(content: str) -> Dict[str, Any]:
    """解析大模型返回：{"project_name": 文件中的项目名或 None, "items": [...]}。"""
    payload = _load_llm_json(content)
    project_name = payload.get("projectName") if isinstance(payload, dict) else None
    project_name = project_name.strip() if isinstance(project_name, str) and project_name.strip() else None
    items = payload.get("items") if isinstance(payload, dict) else None
    return {"project_name": project_name, "items": _normalize_items(items)}


def parse_llm_items(content: str) -> List[Dict[str, Optional[str]]]:
    """仅取归一化 items（兼容旧调用与单测）。"""
    return parse_llm_payload(content)["items"]


# ── 第四步：与现有节点匹配、分桶 ───────────────────────────────────

def _find_child(flat: List[Dict], parent: Optional[Dict], segment: str) -> Optional[Dict]:
    """在 parent 的子节点（parent 为 None 时找根节点）中按标题找节点。"""
    parent_id = parent["id"] if parent else None
    target = _norm(segment)
    return next((n for n in flat if n["parent_id"] == parent_id and _norm(n["title"]) == target), None)


def resolve_parent(flat: List[Dict], suggested_path: Optional[str]) -> Tuple[Optional[str], Optional[str]]:
    """「建议归属路径」→ 现有节点 (id, path)。

    先按路径逐段向下找；整条路径对不上时退化为按最后一段标题全局找（例如大模型只给了
    「订单信息」）。新节点深度 = 归属层 + 1，不得超过 MAX_NODE_DEPTH，超深时上提到
    允许的最深层。找不到任何归属时返回 (None, None)，由前端用「导入信息」兜底。
    """
    if not suggested_path:
        return None, None
    segments = [s.strip() for s in re.split(r"[／/>＞|]+", str(suggested_path)) if s.strip()]
    if not segments:
        return None, None

    deepest: Optional[Dict] = None
    current: Optional[Dict] = None
    for segment in segments:
        child = _find_child(flat, current, segment)
        if child is None:
            break
        current = child
        deepest = child
    if deepest is None:
        deepest = next((n for n in flat if _norm(n["title"]) == _norm(segments[-1])), None)
    if deepest is None:
        return None, None

    by_id = {n["id"]: n for n in flat}
    while deepest["depth"] > MAX_NODE_DEPTH - 1:
        parent = by_id.get(deepest["parent_id"])
        if parent is None:
            return None, None
        deepest = parent
    return deepest["id"], deepest["path"]


def _pick_same_title(candidates: List[Dict], parent_hint: Optional[str]) -> Dict:
    """同名节点（如两个「密码」）先按建议归属路径消歧，消歧不了取第一个。"""
    if len(candidates) == 1 or not parent_hint:
        return candidates[0]
    hint = _norm(parent_hint)
    return next((c for c in candidates if hint and hint in _norm(c["path"])), candidates[0])


def _prefer_value_holder(candidates: List[Dict], parent_hint: Optional[str], value: str) -> Dict:
    """同名节点消歧：建议归属路径 → 能装下这个值的下拉 → 第一个。

    旧版导入把车型写成「标题=型号」的文本节点，项目下于是留下过 text 版「车型1」，
    与全局的「车型1」下拉同名。按标题取第一个常取到那个孤儿节点，值写进文本节点、
    下拉始终空着——所以要优先挑真能装下这个值的那一个。
    """
    picked = _pick_same_title(candidates, parent_hint)
    hint = _norm(parent_hint or "")
    if len(candidates) == 1 or not value or (hint and hint in _norm(picked["path"])):
        return picked
    holder = next((c for c in candidates
                   if c["content_type"] == "select" and snap_select_value(value, c["options"]) is not None), None)
    return holder if holder is not None else picked


def snap_select_value(value: str, options: List[str]) -> Optional[str]:
    """识别值 → 节点可选项（对不上返回 None，由调用方按未匹配处理）。

    先精确比（忽略大小写、空白与常见标点）；不中时，**只有可选项是车型型号**才再放宽一层：
    型号大小写/连字符差异、旧型号名（XS1161→XS1201）、值里夹带中文全称或数量
    （「XCD101 潜伏顶升搬运机器人」「2 台 XCD101」）。一个值里认出多个不同型号时
    返回 None —— 宁可让用户手动归属，也不要蒙一个。
    普通下拉（项目类型等）不放宽：「试点项目一期」不该被吸到「试点项目」上。
    """
    exact = next((option for option in options if _norm(option) == _norm(value)), None)
    if exact is not None:
        return exact
    if not value:
        return None

    by_norm = {_norm(code): code for code in VEHICLE_MODEL_CODES}
    model_options = {_norm(option): option for option in options if _norm(option) in by_norm}
    if not model_options:
        return None

    hits: List[str] = []
    for token in _MODEL_TOKEN_RE.findall(value):
        key = _norm(token)
        key = _norm(VEHICLE_MODEL_ALIASES.get(key.upper(), key))
        code = by_norm.get(key)
        if code is None or _norm(code) not in model_options:
            continue
        option = model_options[_norm(code)]
        if option not in hits:
            hits.append(option)
    return hits[0] if len(hits) == 1 else None


def match_items(flat: List[Dict], items: List[Dict[str, Optional[str]]]) -> Dict[str, List[Dict]]:
    """识别条目 → 现有节点匹配，返回 {fill, overwrite, unmatched} 三类预览数据。"""
    # 候选=能填值的节点。不能只看末级：车型1 是下拉且带「数量」子节点，
    # 按「末级」筛会把它整个排除在匹配之外，车型永远只能当新节点建。
    leaves = [n for n in flat if is_fillable_node(n) and n["content_type"] in ("text", "select")]
    by_title: Dict[str, List[Dict]] = {}
    by_path: Dict[str, Dict] = {}
    for leaf in leaves:
        by_title.setdefault(_norm(leaf["title"]), []).append(leaf)
        by_path.setdefault(_norm(leaf["path"]), leaf)

    fill: List[Dict] = []
    overwrite: List[Dict] = []
    unmatched: List[Dict] = []
    seen_node_ids: set = set()

    for item in items:
        title = item["title"] or ""
        value = item["value"] or ""
        node_title = item["node_title"]
        parent_hint = item["suggested_parent_path"]

        # 1) 精确匹配：节点标题或完整路径（大模型给的 nodeTitle 优先）
        match: Optional[Dict] = None
        for key in (node_title, title):
            if not key:
                continue
            key_norm = _norm(key)
            if key_norm in by_path:
                match = by_path[key_norm]
                break
            if key_norm in by_title:
                match = _prefer_value_holder(by_title[key_norm], parent_hint, value)
                break

        # 2) 相似度兜底：与可填节点标题做序列相似度，达到阈值（0.9）才认
        if match is None:
            best: Optional[Dict] = None
            best_score = 0.0
            for leaf in leaves:
                leaf_norm = _norm(leaf["title"])
                score = _ratio(_norm(title), leaf_norm)
                if node_title:
                    score = max(score, _ratio(_norm(node_title), leaf_norm))
                if score > best_score:
                    best, best_score = leaf, score
            if best is not None and best_score >= SIMILARITY_THRESHOLD:
                match = best

        # 3) 匹配到下拉节点：识别值必须命中可选项（车型型号再放宽一层），否则按未匹配处理
        if match is not None and match["content_type"] == "select":
            canonical = snap_select_value(value, match["options"])
            if canonical is None:
                match = None
            else:
                value = canonical

        if match is None:
            parent_id, parent_path = resolve_parent(flat, parent_hint)
            unmatched.append({
                "title": title or (node_title or ""),
                "value": value,
                "quantity": item.get("quantity") or None,
                "suggested_parent_id": parent_id,
                "suggested_parent_path": parent_path,
            })
            continue

        if match["id"] in seen_node_ids:
            continue
        seen_node_ids.add(match["id"])

        current = _current_text(match)
        if _norm(current) != _norm(value):  # 与现有内容一致时不产生变更
            (overwrite if current.strip() else fill).append(_matched_row(match, value))

        # 车型条目自带的「数量」直接落到该车型的下一个子节点（通常是「数量」）：
        # 不靠大模型再给一条 nodeTitle=数量 的条目——同名子节点在车型1/车型2 下都有，
        # 让它自己说清属于哪辆车太不可靠。落到子节点后进同一个 fill/overwrite 预览。
        quantity = item.get("quantity")
        if quantity and match["content_type"] == "select":
            child = _find_child(flat, match, "数量")
            if child is not None and child["id"] not in seen_node_ids and is_fillable_node(child):
                seen_node_ids.add(child["id"])
                child_current = _current_text(child)
                if _norm(child_current) != _norm(quantity):
                    (overwrite if child_current.strip() else fill).append(_matched_row(child, quantity))
            elif child is None:
                # 该车型下还没有「数量」子节点（只有一个光杆下拉，如 车型3）：数量不能吞掉，
                # 推成一条未匹配，导到车型节点下增补出「数量」——与既有车型的形状保持一致
                unmatched.append({
                    "title": "数量",
                    "value": quantity,
                    "quantity": None,
                    "suggested_parent_id": match["id"],
                    "suggested_parent_path": match["path"],
                })

    return {"fill": fill, "overwrite": overwrite, "unmatched": unmatched}


def _matched_row(node: Dict, value: str) -> Dict:
    """匹配条目 → 预览行（fill / overwrite 共用）。"""
    return {
        "node_id": node["id"],
        "path": node["path"],
        "title": node["title"],
        "content_type": node["content_type"],
        "current": _current_text(node),
        "value": value,
    }


# ── 第五步：调大模型 + 编排 ────────────────────────────────────────

async def _call_llm(prompt: str) -> str:
    """调用「摇人」同款大模型客户端（LLM_API_KEY + LLM_MODEL_NAME，默认 DeepSeek flash）。"""
    from app.modules.call.services.model_service import ModelService

    try:
        client = ModelService.get_client()
    except ValueError as exc:  # LLM_API_KEY 未配置
        raise RuntimeError("AI 服务未配置：请在 backend/.env 设置 LLM_API_KEY（与摇人共用同一密钥）") from exc

    try:
        response = await client.chat.completions.create(
            model=settings.LLM_MODEL_NAME,
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": prompt},
            ],
            temperature=LLM_TEMPERATURE,
            stream=False,
            timeout=LLM_TIMEOUT_SECONDS,
        )
    except Exception as exc:  # 网络/鉴权/超时等统一给用户可读信息
        raise RuntimeError(f"大模型调用失败：{exc}") from exc

    content = response.choices[0].message.content if response.choices else None
    if not content:
        raise RuntimeError("大模型没有返回内容，请稍后重试")
    return content


async def analyze_import_file(project_id: str, filename: str, data: bytes) -> Dict[str, Any]:
    """文件 → 识别 → 匹配，返回三类预览（不落库）。

    异常约定（接口层映射）：ValueError → 400；RuntimeError → 503（AI 未配置/调用失败）。
    """
    if not data:
        raise ValueError("文件内容为空")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(f"文件不能超过 {MAX_FILE_BYTES // (1024 * 1024)}MB，请拆分后再导入")

    # docx 解压 / xlsx 解析是 CPU 密集的同步代码，放线程池执行，
    # 避免大文件解析期间阻塞事件循环、冻结其他并发请求
    text = await asyncio.to_thread(extract_text, filename, data)
    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS]

    roots = info_node_service.get_tree(project_id)
    if not roots:
        raise ValueError("该项目还没有信息节点，请先在编辑页初始化信息树或新建节点后再导入")

    project_name = _get_project_name(project_id)
    flat = flatten_tree(roots)
    # 信息树有车辆/车型节点时注入 AGV 车型清单，让车型信息的 title 直接用车型型号
    vehicle_parent = find_vehicle_parent_path(flat)
    prompt = build_import_prompt(build_node_catalog(flat), text, project_name, vehicle_parent)
    content = await _call_llm(prompt)
    parsed = parse_llm_payload(content)
    items = parsed["items"]
    buckets = match_items(flat, items)

    return {
        "file_name": filename,
        "model": settings.LLM_MODEL_NAME,
        "text_length": len(text),
        "truncated": truncated,
        "extracted": len(items),
        # 项目名校对：文件里识别到的项目名与系统内不一致时，前端提醒可能导错文件
        "project_name": project_name,
        "file_project_name": parsed["project_name"],
        "name_mismatch": project_name_mismatch(project_name, parsed["project_name"]),
        **buckets,
    }
